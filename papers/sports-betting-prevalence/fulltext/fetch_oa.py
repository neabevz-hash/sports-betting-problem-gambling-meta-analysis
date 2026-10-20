"""Open-access full texts for the reports sent to full-text screening (protocol, section 11b: open access only, no author contact).

Only legal open-access sources, no author contact. For every record with final_decision include or unsure in
screening/decisions/tiab_final.csv, in this order:
  1. Europe PMC full-text XML (PMC open-access subset), found by PMID or DOI: text with sections and tables;
  2. CORE full text (API v3 with CORE_API_KEY), by the record's CORE ids, then by DOI;
  3. PDFs from open-access links: Unpaywall, OpenAlex, the record's own link (CORE, OpenAIRE, Scholar), and the
     citation_pdf_url tag of the DOI landing page (repositories).
A page that answers with a bot check is recorded as blocked and not worked around.

Output (git-ignored, publishers' texts are not redistributed): fulltext/text/<rec_id>.txt with page or section
markers, fulltext/pdf/<rec_id>.pdf. Committed: fulltext/retrieval_log.csv (what was tried, what worked, SHA-256).
Resumable: records already retrieved are skipped.

Template: papers/misophonia-prevalence/fulltext/fetch_oa.py (unchanged logic).
Run: .venv/Scripts/python.exe -I papers/sports-betting-prevalence/fulltext/fetch_oa.py --run 2026-10-18 [--no-core]
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import html
import json
import re
import sys
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path

import httpx
from lxml import etree

HERE = Path(__file__).resolve().parent
PAPER = HERE.parent
ROOT = PAPER.parents[1]
TEXT, PDF = HERE / "text", HERE / "pdf"
LOG = HERE / "retrieval_log.csv"
LOG_FIELDS = ["rec_id", "doi", "status", "source", "url", "chars", "pages", "sha256", "tried", "at_utc"]
EMAIL = "mrexert@gmail.com"
BOT = re.compile(r"just a moment|cf-chl|captcha|are you a robot|access denied|unusual traffic|making sure you", re.I)
MAX_PDF = 60 * 1024 * 1024


def env() -> dict[str, str]:
    out = {}
    for line in (ROOT / ".env").read_text(encoding="utf-8").splitlines():
        if "=" in line and not line.strip().startswith("#"):
            k, v = line.split("=", 1)
            out[k.strip()] = v.strip()
    return out


ENV = env()
client = httpx.Client(timeout=httpx.Timeout(30, connect=15), follow_redirects=True,
                      limits=httpx.Limits(max_connections=32, max_keepalive_connections=16),
                      headers={"User-Agent": f"sports-betting-prevalence-fulltext/1.0 (mailto:{EMAIL})"})


def utcnow() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def get(url: str, **kw) -> httpx.Response | None:
    for attempt in range(3):
        try:
            r = client.get(url, **kw)
        except httpx.HTTPError:
            time.sleep(3 * (attempt + 1))
            continue
        if r.status_code in (429, 500, 502, 503, 504):
            time.sleep(5 * (attempt + 1))
            continue
        return r
    return None


# ------------------------------------------------------------------ 1. Europe PMC XML
def jats_to_text(xml: bytes) -> str:
    root = etree.fromstring(xml, parser=etree.XMLParser(recover=True, huge_tree=True))
    txt = lambda el: " ".join("".join(el.itertext()).split())  # noqa: E731
    out = []
    t = root.find(".//article-meta/title-group/article-title")
    if t is not None:
        out.append(f"TITLE: {txt(t)}")
    ab = root.find(".//article-meta/abstract")
    if ab is not None:
        out.append(f"=== Abstract ===\n{txt(ab)}")
    body = root.find(".//body")
    if body is not None:
        for el in body.iter("sec", "p", "table-wrap"):
            if el.tag == "sec":
                title = el.find("title")
                if title is not None:
                    out.append(f"=== Section: {txt(title)} ===")
            elif el.tag == "p" and el.getparent() is not None and el.getparent().tag != "caption":
                if not any(a.tag == "table-wrap" for a in el.iterancestors()):
                    out.append(txt(el))
            elif el.tag == "table-wrap":
                label = el.find("label")
                cap = el.find("caption")
                out.append(f"=== {txt(label) if label is not None else 'Table'}: {txt(cap) if cap is not None else ''} ===")
                for tr in el.iter("tr"):
                    out.append(" | ".join(txt(c) for c in tr if c.tag in ("td", "th")))
                foot = el.find("table-wrap-foot")
                if foot is not None:
                    out.append(f"[table note] {txt(foot)}")
    return "\n".join(x for x in out if x)


def europepmc(rec: dict) -> tuple[str | None, str, str]:
    q = f"EXT_ID:{rec['pmid']} AND SRC:MED" if rec.get("pmid") else (f'DOI:"{rec["doi"]}"' if rec.get("doi") else None)
    if not q:
        return None, "", "epmc:no id"
    r = get("https://www.ebi.ac.uk/europepmc/webservices/rest/search", params={"query": q, "format": "json", "resultType": "core"})
    hits = (r.json().get("resultList") or {}).get("result", []) if r is not None and r.status_code == 200 else []
    pmcid = next((h.get("pmcid") for h in hits if h.get("pmcid")), None)
    if not pmcid:
        return None, "", "epmc:no pmcid"
    rec["_pmcid"] = pmcid   # for the rendered-PDF fallback
    url = f"https://www.ebi.ac.uk/europepmc/webservices/rest/{pmcid}/fullTextXML"
    x = get(url)
    if x is None or x.status_code != 200 or not x.content.strip().startswith(b"<"):
        return None, url, f"epmc:{pmcid} no xml"
    text = jats_to_text(x.content)
    return (text, url, f"epmc:{pmcid} ok") if len(text) > 1500 else (None, url, f"epmc:{pmcid} short")


# ------------------------------------------------------------------ 2. CORE full text
def core_get(url: str, params: dict | None = None) -> httpx.Response | None:
    for attempt in range(6):
        r = get(url, params=params, headers={"Authorization": f"Bearer {ENV.get('CORE_API_KEY', '')}"})
        if r is not None and r.status_code == 429:
            time.sleep(30 * (attempt + 1))
            continue
        time.sleep(0.5)
        return r
    return None


def core(rec: dict) -> tuple[str | None, str, str]:
    if not ENV.get("CORE_API_KEY"):
        return None, "", "core:no key"
    ids = [s.split(":", 1)[1] for s in rec["source_ids"].split(";") if s.startswith("CORE:")]
    tried = []
    for cid in ids[:3]:
        url = f"https://api.core.ac.uk/v3/works/{cid}"
        r = core_get(url)
        ft = (r.json().get("fullText") or "") if r is not None and r.status_code == 200 else ""
        tried.append(f"core:{cid} {'ok' if len(ft) > 1500 else 'none'}")
        if len(ft) > 1500:
            return ft, url, "; ".join(tried)
    if rec.get("doi"):
        r = core_get("https://api.core.ac.uk/v3/search/works/", {"q": f'doi:"{rec["doi"]}"', "limit": 3})
        for w in (r.json().get("results") or []) if r is not None and r.status_code == 200 else []:
            ft = w.get("fullText") or ""
            if len(ft) > 1500:
                tried.append(f"core:doi ok ({w.get('id')})")
                return ft, f"https://core.ac.uk/works/{w.get('id')}", "; ".join(tried)
        tried.append("core:doi none")
    else:   # theses and other records without a DOI: exact title search
        title = re.sub(r'["\\]', " ", rec.get("title") or "")[:150].strip()
        if len(title) >= 25:
            r = core_get("https://api.core.ac.uk/v3/search/works/", {"q": f'title:"{title}"', "limit": 3})
            for w in (r.json().get("results") or []) if r is not None and r.status_code == 200 else []:
                ft = w.get("fullText") or ""
                if len(ft) > 1500 and compact(w.get("title")) == compact(rec.get("title")):
                    tried.append(f"core:title ok ({w.get('id')})")
                    return ft, f"https://core.ac.uk/works/{w.get('id')}", "; ".join(tried)
            tried.append("core:title none")
    return None, "", "; ".join(tried) or "core:no id"


def compact(t: str | None) -> str:
    import unicodedata
    t = unicodedata.normalize("NFKD", re.sub(r"<[^>]+>", " ", t or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "", t)


S2_LOCK = threading.Lock()
USE_S2 = False   # off: without a key Semantic Scholar answers 429 and stalls the run (2026-10-11)


def s2_pdf(rec: dict) -> list[tuple[str, str]]:
    """Semantic Scholar openAccessPdf (some PDFs are served from Semantic Scholar's own servers)."""
    ids = []
    if rec.get("doi"):
        ids.append(f"DOI:{rec['doi']}")
    ids += [s.split(":", 1)[1] for s in rec["source_ids"].split(";") if s.startswith("S2:")]
    out = []
    for pid in ids[:2]:
        with S2_LOCK:   # about one request a second without a key
            r = get(f"https://api.semanticscholar.org/graph/v1/paper/{pid}", params={"fields": "openAccessPdf"})
            time.sleep(1.1)
        if r is not None and r.status_code == 200:
            u = ((r.json() or {}).get("openAccessPdf") or {}).get("url")
            if u:
                out.append(("semanticscholar", u))
                break
    return out


def repository_pdfs(doi: str) -> list[tuple[str, str]]:
    """Direct file links from repository APIs: Zenodo, figshare (also institutional figshare), OSF preprints."""
    out = []
    d = re.sub(r"([./_])v\d+$", "", doi.lower())
    m = re.match(r"10\.5281/zenodo\.(\d+)$", d)
    if m:
        r = get(f"https://zenodo.org/api/records/{m.group(1)}")
        if r is None or r.status_code != 200:   # a concept DOI: take its latest version
            r = get(f"https://zenodo.org/api/records/{m.group(1)}/versions/latest")
        if r is not None and r.status_code == 200:
            for f in r.json().get("files") or []:
                link = (f.get("links") or {}).get("self") or ""
                if (f.get("key") or "").lower().endswith(".pdf") and link:
                    out.append(("zenodo", link))
    fig = r"10\.(?:6084/m9\.figshare\.|25959/|25949/|25405/ncl\.|1184/r1/|26180/|26188/|17608/k6\.auckland\.)"
    m = re.match(fig + r"(\d+)$", d)
    if m:
        out += figshare_article_pdfs(m.group(1))
    m = re.match(fig + r"c\.(\d+)$", d)   # a figshare collection: the PDFs of its articles
    if m:
        r = get(f"https://api.figshare.com/v2/collections/{m.group(1)}/articles", params={"page_size": 50})
        if r is not None and r.status_code == 200:
            for art in r.json()[:20]:
                out += figshare_article_pdfs(str(art.get("id")))
    m = re.match(r"10\.(?:31234|31219|17605|31235|31730)/osf\.io/([a-z0-9]+)", d)
    if m:
        out += osf_pdfs(m.group(1))
    return out


def figshare_article_pdfs(aid: str) -> list[tuple[str, str]]:
    r = get(f"https://api.figshare.com/v2/articles/{aid}")
    if r is None or r.status_code != 200:
        return []
    return [("figshare", f["download_url"]) for f in r.json().get("files") or []
            if (f.get("name") or "").lower().endswith(".pdf") and f.get("download_url")]


def osf_pdfs(guid: str) -> list[tuple[str, str]]:
    """OSF through its API: a preprint's primary file, or the PDFs stored in a project or registration."""
    out = []
    r = get(f"https://api.osf.io/v2/preprints/{guid}/")
    if r is not None and r.status_code == 200:
        rel = ((r.json().get("data") or {}).get("relationships") or {}).get("primary_file") or {}
        href = ((rel.get("links") or {}).get("related") or {}).get("href")
        f = get(href) if href else None
        if f is not None and f.status_code == 200:
            dl = (((f.json().get("data") or {}).get("links")) or {}).get("download")
            if dl:
                out.append(("osf_preprint", dl))
    for kind in ("nodes", "registrations"):
        r = get(f"https://api.osf.io/v2/{kind}/{guid}/files/osfstorage/")
        if r is not None and r.status_code == 200:
            items = list((r.json().get("data") or [])[:30])
            for item in [i for i in items if (i.get("attributes") or {}).get("kind") == "folder"][:5]:   # one level down
                sub = (((item.get("relationships") or {}).get("files") or {}).get("links") or {}).get("related", {}).get("href")
                s = get(sub) if sub else None
                if s is not None and s.status_code == 200:
                    items += (s.json().get("data") or [])[:30]
            for item in items:
                a = item.get("attributes") or {}
                if a.get("kind") == "file" and (a.get("name") or "").lower().endswith(".pdf"):
                    dl = (item.get("links") or {}).get("download")
                    if dl:
                        out.append(("osf_file", dl))
            break
    out.append(("osf", f"https://osf.io/{guid}/download"))
    return out


SCHOLAR_LINKS: list[tuple[str, str]] = []


def load_scholar_links(run: str) -> None:
    """[PDF] side links of the Google Scholar result pages R1 saved (kept locally only)."""
    for f in sorted((PAPER / "search" / "runs" / run / "manual").glob("scholar_*.html")):
        t = f.read_text(encoding="utf-8", errors="replace")
        for chunk in re.split(r'<div class="gs_r gs_or gs_scl"', t)[1:]:
            side = re.search(r'<div class="gs_ggs gs_fl">.*?href="([^"]+)"', chunk, re.S)
            h3 = re.search(r'<h3 class="gs_rt"[^>]*>(.*?)</h3>', chunk, re.S)
            if side and h3:
                frag = max(re.split(r"…|\.\.\.", re.sub(r"<[^>]+>", "", h3.group(1))), key=len)
                key = compact(re.sub(r"^\s*(\[[A-Z]+\]\s*)+", "", frag))
                if len(key) >= 25:
                    SCHOLAR_LINKS.append((key, html.unescape(side.group(1))))


def scholar_pdf(rec: dict) -> list[tuple[str, str]]:
    t = compact(rec.get("title"))
    return [("scholar_side_link", u) for key, u in SCHOLAR_LINKS if key in t]


# ------------------------------------------------------------------ 3. PDFs
def pdf_candidates(rec: dict) -> list[tuple[str, str]]:
    cands = []
    doi = rec.get("doi")
    if doi:
        r = get(f"https://api.unpaywall.org/v2/{doi}", params={"email": EMAIL})
        if r is not None and r.status_code == 200:
            for loc in r.json().get("oa_locations") or []:
                for key in ("url_for_pdf", "url"):
                    if loc.get(key):
                        cands.append(("unpaywall", loc[key]))
        r = get(f"https://api.openalex.org/works/doi:{doi}",
                params={"select": "best_oa_location,locations", "api_key": ENV.get("OPENALEX_API_KEY", "")})
        if r is not None and r.status_code == 200:
            w = r.json()
            for loc in [w.get("best_oa_location") or {}] + (w.get("locations") or []):
                if loc and loc.get("is_oa") and loc.get("pdf_url"):
                    cands.append(("openalex", loc["pdf_url"]))
    if doi:
        cands += repository_pdfs(doi)
    pmcid = rec.get("_pmcid")
    if pmcid:
        cands.append(("europepmc_render", f"https://europepmc.org/articles/{pmcid}?pdf=render"))
    if USE_S2:
        cands += s2_pdf(rec)
    cands += scholar_pdf(rec)
    if rec.get("url"):
        cands.append(("record", rec["url"]))
    if doi:
        cands.append(("landing", f"https://doi.org/{doi}"))
    seen, out = set(), []
    for src, u in cands:
        if u and u not in seen:
            seen.add(u)
            out.append((src, u))
    return out


def fetch_pdf(url: str) -> tuple[bytes | None, str]:
    r = get(url)
    if r is None:
        return None, "network"
    body = r.content[:MAX_PDF]
    if body[:5] == b"%PDF-":
        return body, "pdf"
    head = body[:20000].decode("utf-8", "ignore")
    if BOT.search(head):
        return None, "blocked"
    m = re.search(r'<meta[^>]+name=["\']citation_pdf_url["\'][^>]+content=["\']([^"\']+)', head, re.I) \
        or re.search(r'<meta[^>]+content=["\']([^"\']+)["\'][^>]+name=["\']citation_pdf_url', head, re.I)
    if m:
        r2 = get(m.group(1))
        if r2 is not None and r2.content[:5] == b"%PDF-":
            return r2.content[:MAX_PDF], "pdf via citation_pdf_url"
        if r2 is not None and BOT.search(r2.content[:20000].decode("utf-8", "ignore")):
            return None, "blocked"
    return None, f"not pdf ({r.status_code})"


def pdf_to_text(data: bytes) -> tuple[str, int]:
    import pymupdf
    doc = pymupdf.open(stream=data, filetype="pdf")
    parts = [f"=== Page {i + 1} ===\n{page.get_text()}" for i, page in enumerate(doc)]
    return "\n".join(parts), doc.page_count


def pdfs(rec: dict) -> tuple[str | None, str, str, int, bytes | None]:
    tried = []
    for src, url in pdf_candidates(rec):
        data, how = fetch_pdf(url)
        tried.append(f"{src}:{how}")
        if data:
            try:
                text, pages = pdf_to_text(data)
            except Exception as e:  # noqa: BLE001 - a broken PDF is just a failed candidate
                tried[-1] = f"{src}:unreadable pdf ({type(e).__name__})"
                continue
            if len(text) > 1500:
                return text, url, "; ".join(tried), pages, data
            tried[-1] = f"{src}:pdf without text"
        time.sleep(0.5)
    return None, "", "; ".join(tried) or "pdf:no candidates", 0, None


# ------------------------------------------------------------------ main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--run", required=True)
    ap.add_argument("--limit", type=int)
    ap.add_argument("--workers", type=int, default=8)
    ap.add_argument("--import-manual", action="store_true", help="take PDFs R1 saved into fulltext/manual/")
    ap.add_argument("--all-records", action="store_true", help="every record of the run goes to full text")
    ap.add_argument("--no-core", action="store_true",
                    help="skip CORE (its API timed out on 2026-10-19; records without a text get a CORE retry later, D2)")
    ap.add_argument("--only", nargs="*", help="retry only these records")
    a = ap.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")
    TEXT.mkdir(parents=True, exist_ok=True)
    PDF.mkdir(parents=True, exist_ok=True)
    recs = {}
    for line in (PAPER / "search" / "runs" / a.run / "records_unique.jsonl").open(encoding="utf-8"):
        u = json.loads(line)
        recs[u["rec_id"]] = u
    if a.all_records:   # a run whose records all go to full text (citation search, chase_records.py)
        final = [{"rec_id": k} for k in sorted(recs)]
    else:
        final = [r for r in csv.DictReader((PAPER / "screening" / "decisions" / "tiab_final.csv").open(encoding="utf-8-sig"))
                 if r["full_text"] == "True"]
    done = {}
    if LOG.exists():
        done = {r["rec_id"]: r for r in csv.DictReader(LOG.open(encoding="utf-8-sig"))}
    todo = [r["rec_id"] for r in final if done.get(r["rec_id"], {}).get("status") != "retrieved"]
    already = len(final) - len(todo)
    if a.import_manual:   # PDFs R1 saved from a browser into fulltext/manual/<rec_id>.pdf (MANUAL_PDFS.md)
        added = 0
        for p in sorted([*(HERE / "manual").glob("R*.pdf"), *(HERE / "manual").glob("R*.docx")]):
            rid = p.stem
            if rid not in todo:
                continue
            data = p.read_bytes()
            if p.suffix.lower() == ".docx":   # an author manuscript in Word: paragraphs and table rows
                from docx import Document
                d = Document(str(p))
                rows = [" | ".join(c.text for c in r.cells) for t in d.tables for r in t.rows]
                text, pages = "\n".join([par.text for par in d.paragraphs] + ["=== Tables ==="] + rows), ""
            elif data[:5] != b"%PDF-":
                print(f"{p.name}: not a PDF, skipped")
                continue
            else:
                text, pages = pdf_to_text(data)
            now = utcnow()
            kind = "manual_docx" if p.suffix.lower() == ".docx" else "manual_pdf"
            (TEXT / f"{rid}.txt").write_text(f"REC_ID: {rid}\nSOURCE: {kind}\nURL: saved by R1 from a browser\n"
                                             f"RETRIEVED_UTC: {now}\n\n{text}", encoding="utf-8")
            (PDF / f"{rid}{p.suffix.lower()}").write_bytes(data)
            prev = done.get(rid, {})
            done[rid] = {**{k: prev.get(k, "") for k in LOG_FIELDS}, "rec_id": rid, "status": "retrieved",
                         "source": kind, "url": "saved by R1 from a browser", "chars": len(text), "pages": pages,
                         "sha256": hashlib.sha256(data).hexdigest(), "at_utc": now,
                         "tried": (prev.get("tried", "") + f" | {kind} ok").strip(" |")}
            added += 1
        with LOG.open("w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=LOG_FIELDS)
            w.writeheader()
            w.writerows(done[k] for k in sorted(done))
        print(f"imported {added} PDFs saved by R1")
        return
    if a.only:
        todo = [r for r in todo if r in set(a.only)]
    if a.limit:
        todo = todo[:a.limit]
    print(f"{len(final)} reports, {already} already retrieved, {len(todo)} to try", flush=True)
    load_scholar_links(a.run)
    print(f"{len(SCHOLAR_LINKS)} [PDF] links from the saved Google Scholar pages", flush=True)

    def process(rid: str) -> dict:
        rec = recs[rid]
        row = {"rec_id": rid, "doi": rec.get("doi", ""), "status": "not retrieved", "source": "", "url": "",
               "chars": 0, "pages": "", "sha256": "", "tried": "", "at_utc": utcnow()}
        tried = []
        text, url, note = europepmc(rec)
        tried.append(note)
        source, pages, data = "europepmc_xml", "", None
        if text is None and not a.no_core:
            text, url, note = core(rec)
            tried.append(note)
        elif text is None:
            tried.append("core:skipped (API down)")
            source = "core_fulltext"
        if text is None:
            text, url, note, pages, data = pdfs(rec)
            tried.append(note)
            source = "pdf"
        if text is not None:
            header = f"REC_ID: {rid}\nSOURCE: {source}\nURL: {url}\nRETRIEVED_UTC: {row['at_utc']}\n\n"
            (TEXT / f"{rid}.txt").write_text(header + text, encoding="utf-8")
            if data:
                (PDF / f"{rid}.pdf").write_bytes(data)
            digest = hashlib.sha256(data if data else text.encode("utf-8")).hexdigest()
            row.update(status="retrieved", source=source, url=url, chars=len(text), pages=pages, sha256=digest)
        row["tried"] = " | ".join(t for t in tried if t)
        return row

    lock = threading.Lock()
    with ThreadPoolExecutor(max_workers=a.workers) as pool:
        futures = {pool.submit(process, rid): rid for rid in todo}
        for n, fut in enumerate(as_completed(futures), 1):
            rid = futures[fut]
            try:
                row = fut.result()
            except Exception as e:  # noqa: BLE001 - one failed record must not stop the run
                row = {"rec_id": rid, "doi": recs[rid].get("doi", ""), "status": "not retrieved", "source": "",
                       "url": "", "chars": 0, "pages": "", "sha256": "", "tried": f"error: {e!r}"[:300], "at_utc": utcnow()}
            with lock:
                done[rid] = row
                with LOG.open("w", encoding="utf-8-sig", newline="") as f:
                    w = csv.DictWriter(f, fieldnames=LOG_FIELDS)
                    w.writeheader()
                    w.writerows(done[k] for k in sorted(done))
            print(f"[{n}/{len(todo)}] {rid} {row['status']} {row['source']}", flush=True)
    got = sum(1 for r in done.values() if r["status"] == "retrieved")
    print(f"retrieved {got} of {len(final)}", flush=True)


if __name__ == "__main__":
    main()
